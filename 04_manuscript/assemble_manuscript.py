#!/usr/bin/env python3
"""
Manuscript Assembler - combines all section files into a single complete manuscript.
Generates both .md and .docx output.
"""

import logging
from datetime import datetime
from pathlib import Path

logger = logging.getLogger(__name__)

SECTIONS_DIR = Path(__file__).parent / "sections"
OUTPUT_DIR = Path(__file__).parent / "output"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# Section order
SECTION_FILES = [
    "abstract.md",
    "introduction.md",
    "methods.md",
    "results.md",
    "discussion.md",
    "conclusions.md",
]

HEADER = """---
title: "Large Language Models in Healthcare Simulation: A Bibliometric Analysis Mapping the Research Landscape"
authors:
  - Matt Pears (First Author)
  - "[Colleague 1]"
  - "[Colleague 2]"
  - "Shekhar [Surname] (Corresponding Author)"
date: {date}
---

"""

FOOTER = """
---

## Data Availability Statement

The complete data collection pipeline, processing scripts, and analysis code are available at https://github.com/MattPears1/bibliometric-llm-healthcare-simulation. Raw bibliometric data and analysis outputs are available upon reasonable request from the corresponding author.

## Funding

[To be completed]

## Competing Interests

The authors declare no competing interests.

## Author Contributions (CRediT)

- **Matt Pears:** Conceptualisation, Methodology, Software, Formal Analysis, Data Curation, Writing -- Original Draft, Visualisation
- **[Colleague 1]:** Validation (inter-rater reliability), Writing -- Review & Editing
- **[Colleague 2]:** Validation (inter-rater reliability), Writing -- Review & Editing
- **Shekhar [Surname]:** Conceptualisation, Supervision, Writing -- Review & Editing, Project Administration

## Acknowledgements

[To be completed]

---

## Supplementary Materials

- **Appendix A:** Full search strategies per database
- **Appendix B:** PRISMA 2020 checklist
- **Appendix C:** Complete list of included papers
- **Appendix D:** Code availability and reproduction guide
"""


def assemble_markdown():
    """Assemble all sections into a single markdown file."""
    timestamp = datetime.now().strftime("%Y%m%d")
    output_path = OUTPUT_DIR / f"MANUSCRIPT_{timestamp}.md"

    content = HEADER.format(date=datetime.now().strftime("%Y-%m-%d"))

    for section_file in SECTION_FILES:
        section_path = SECTIONS_DIR / section_file
        if section_path.exists():
            with open(section_path, "r", encoding="utf-8") as f:
                section_content = f.read().strip()
            content += section_content + "\n\n"
            logger.info(f"Added section: {section_file}")
        else:
            content += f"\n\n[SECTION MISSING: {section_file}]\n\n"
            logger.warning(f"Missing section: {section_file}")

    content += FOOTER

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(content)

    # Word count
    words = len(content.split())
    logger.info(f"Manuscript assembled: {output_path} ({words} words)")

    return output_path, words


def assemble_docx(md_path: Path):
    """Convert markdown manuscript to .docx using python-docx."""
    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        logger.warning("python-docx not installed. Skipping .docx generation.")
        return None

    doc = Document()

    # Read markdown
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    # Simple markdown to docx conversion
    in_frontmatter = False
    for line in lines:
        line = line.rstrip()

        # Skip YAML frontmatter
        if line == "---":
            in_frontmatter = not in_frontmatter
            continue
        if in_frontmatter:
            continue

        # Headers
        if line.startswith("# "):
            p = doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            p = doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            p = doc.add_heading(line[4:], level=3)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("1. ") or line.startswith("2. ") or line.startswith("3. ") or line.startswith("4. "):
            doc.add_paragraph(line[3:], style="List Number")
        elif line.strip() == "":
            continue
        elif line.startswith("|"):
            # Simple table row - add as paragraph for now
            doc.add_paragraph(line, style="Normal")
        elif line.startswith("**") and line.endswith("**"):
            p = doc.add_paragraph()
            run = p.add_run(line.strip("*"))
            run.bold = True
        else:
            doc.add_paragraph(line)

    docx_path = md_path.with_suffix(".docx")
    doc.save(str(docx_path))
    logger.info(f"DOCX generated: {docx_path}")
    return docx_path


def main():
    logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")

    md_path, words = assemble_markdown()
    print(f"Manuscript assembled: {md_path} ({words} words)")

    docx_path = assemble_docx(md_path)
    if docx_path:
        print(f"DOCX generated: {docx_path}")

    # List what's included and missing
    print("\nSections status:")
    for sf in SECTION_FILES:
        path = SECTIONS_DIR / sf
        if path.exists():
            with open(path, "r") as f:
                wc = len(f.read().split())
            print(f"  [DONE] {sf} ({wc} words)")
        else:
            print(f"  [MISSING] {sf}")


if __name__ == "__main__":
    main()
